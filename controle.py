import re
import unicodedata

import openpyxl

from app.config import PASTA_DADOS


ARQUIVO_CONTROLE = PASTA_DADOS / "CONTROLE_CADASTRO_POSSE.xlsx"
NOME_ABA = "BD"


def normalizar_texto(valor):
    if valor is None:
        return ""

    texto = str(valor).strip().upper()
    texto = unicodedata.normalize("NFD", texto)

    texto = "".join(
        c for c in texto
        if unicodedata.category(c) != "Mn"
    )

    return re.sub(r"[^A-Z0-9]", "", texto)


def normalizar_edital(valor):
    """
    Faz 007/2022 e 7/2022 serem considerados equivalentes.
    """
    if valor is None:
        return ""

    texto = str(valor).strip()

    resultado = re.fullmatch(
        r"0*(\d+)\s*/\s*(\d{4})",
        texto
    )

    if resultado:
        return f"{int(resultado.group(1))}/{resultado.group(2)}"

    return normalizar_texto(texto)


def carregar_registros_por_edital(edital):
    if not ARQUIVO_CONTROLE.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {ARQUIVO_CONTROLE}"
        )

    if not str(edital or "").strip():
        raise ValueError("O edital não foi informado.")

    workbook = openpyxl.load_workbook(
        ARQUIVO_CONTROLE,
        read_only=True,
        data_only=True
    )

    try:
        if NOME_ABA not in workbook.sheetnames:
            raise ValueError(
                f"Aba '{NOME_ABA}' não encontrada. "
                f"Abas disponíveis: {workbook.sheetnames}"
            )

        planilha = workbook[NOME_ABA]
        linhas = planilha.iter_rows(values_only=True)
        cabecalhos = next(linhas, None)

        if not cabecalhos:
            raise ValueError("A planilha não possui cabeçalhos.")

        indices = {
            str(c).strip(): i
            for i, c in enumerate(cabecalhos)
            if c is not None
        }

        indices_normalizados = {
            normalizar_texto(c): c
            for c in indices
        }

        obrigatorias = [
            "EDITAL DE RESULTADO FINAL",
            "CARGO/EMPREGO",
            "ESPECIALIDADE",
            "REGIÃO",
            "NOME",
            "CPF",
        ]

        colunas = {}

        for nome in obrigatorias:
            cabecalho_real = indices_normalizados.get(
                normalizar_texto(nome)
            )

            if cabecalho_real is not None:
                colunas[nome] = cabecalho_real

        faltantes = [
            nome for nome in obrigatorias
            if nome not in colunas
        ]

        if faltantes:
            raise ValueError(
                f"Colunas obrigatórias não encontradas: {faltantes}\n"
                f"Cabeçalhos encontrados: {list(indices)}"
            )

        indice_edital = indices[colunas["EDITAL DE RESULTADO FINAL"]]
        edital_buscado = normalizar_edital(edital)

        print(f"\nEdital recebido: {edital!r}")
        print(f"Edital normalizado: {edital_buscado!r}")
        print(f"Planilha utilizada: {ARQUIVO_CONTROLE}")
        print(f"Aba utilizada: {NOME_ABA}")

        registros = []

        for numero_linha, linha in enumerate(linhas, start=2):
            valor_edital = linha[indice_edital]
            edital_linha = normalizar_edital(valor_edital)

            if edital_linha != edital_buscado:
                continue

            # Preserva todas as colunas da linha original.
            registro = {"LINHA_PLANILHA": numero_linha}

            for cabecalho, indice in indices.items():
                registro[cabecalho] = linha[indice]

            # Garante as chaves padronizadas utilizadas na automação.
            for nome_padrao, cabecalho_real in colunas.items():
                registro[nome_padrao] = registro[cabecalho_real]

            # Associa variações dos nomes dos campos da Página 9.
            aliases = {
                "RG": ["RG"],
                "ORGAO_EXPEDITOR": [
                    "ORGAO EXPEDITOR",
                    "ORGAO EXPEDIDOR",
                ],
                "PCD": ["PCD"],
                "NOME DA MÃE": [
                    "NOME DA MÃE",
                    "NOME DA MAE",
                ],
                "SEXO": ["SEXO"],
                "NASCIMENTO": [
                    "NASCIMENTO",
                    "DATA DE NASCIMENTO",
                ],
                "ESTADO": ["ESTADO", "UF"],
            }

            for destino, alternativas in aliases.items():
                for alternativa in alternativas:
                    origem = indices_normalizados.get(
                        normalizar_texto(alternativa)
                    )

                    if origem is not None:
                        registro[destino] = registro[origem]
                        break

            registros.append(registro)

        if not registros:
            # Diagnóstico para não ocultar o formato real dos editais.
            exemplos = []

            
            for linha in planilha.iter_rows(
                min_row=2,
                max_row=min(planilha.max_row, 20),
                values_only=True
            ):
                valor = linha[indice_edital]
                exemplos.append(
                    f"Linha {len(exemplos) + 2}: {valor!r}"
                )

            print(
                "Cabeçalho da coluna de edital:",
                colunas["EDITAL DE RESULTADO FINAL"]
            )
            print(
                "Índice da coluna de edital:",
                indice_edital
            )
            print(
                "Total de linhas na planilha:",
                planilha.max_row
            )


            raise ValueError(
                f"Nenhum registro encontrado para {edital!r}.\n"
                f"Valor normalizado: {edital_buscado!r}\n"
                f"Exemplos de valores na coluna de edital: {exemplos}"
            )

        print(f"Registros encontrados: {len(registros)}")
        return registros

    finally:
        workbook.close()


def atualizar_status_pronto(registro, status):
    """
    Atualiza somente a coluna PRONTO do candidato identificado
    por LINHA_PLANILHA.

    status deve ser 'OK' ou 'NÃO'.
    """

    if status not in {"OK", "NÃO"}:
        raise ValueError(
            f"Status inválido: {status!r}. Use 'OK' ou 'NÃO'."
        )

    numero_linha = registro.get("LINHA_PLANILHA")

    if not isinstance(numero_linha, int) or numero_linha < 2:
        raise ValueError(
            f"Linha da planilha inválida: {numero_linha!r}"
        )

    if not ARQUIVO_CONTROLE.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {ARQUIVO_CONTROLE}"
        )

    workbook = openpyxl.load_workbook(ARQUIVO_CONTROLE)

    try:
        if NOME_ABA not in workbook.sheetnames:
            raise ValueError(
                f"Aba '{NOME_ABA}' não encontrada na planilha."
            )

        planilha = workbook[NOME_ABA]

        # Localiza a coluna PRONTO sem depender da posição.
        coluna_pronto = None

        for coluna in range(1, planilha.max_column + 1):
            cabecalho = planilha.cell(
                row=1,
                column=coluna
            ).value

            if normalizar_texto(cabecalho) == "PRONTO":
                coluna_pronto = coluna
                break

        if coluna_pronto is None:
            raise ValueError(
                "A coluna 'PRONTO' não foi encontrada "
                f"na aba '{NOME_ABA}'."
            )

        # Confere se a linha ainda corresponde ao CPF esperado.
        coluna_cpf = None

        for coluna in range(1, planilha.max_column + 1):
            cabecalho = planilha.cell(
                row=1,
                column=coluna
            ).value

            if normalizar_texto(cabecalho) == "CPF":
                coluna_cpf = coluna
                break

        if coluna_cpf is None:
            raise ValueError(
                "A coluna CPF não foi encontrada na planilha."
            )

        cpf_planilha = re.sub(
            r"\D",
            "",
            str(
                planilha.cell(
                    row=numero_linha,
                    column=coluna_cpf
                ).value or ""
            )
        )

        cpf_registro = re.sub(
            r"\D",
            "",
            str(registro.get("CPF") or "")
        )

        if cpf_planilha != cpf_registro:
            raise RuntimeError(
                "O CPF da linha da planilha não corresponde "
                "ao candidato que está sendo processado. "
                f"Linha: {numero_linha}; "
                f"CPF esperado: {cpf_registro}; "
                f"CPF encontrado: {cpf_planilha}. "
                "O status não foi alterado."
            )

        planilha.cell(
            row=numero_linha,
            column=coluna_pronto
        ).value = status

        workbook.save(ARQUIVO_CONTROLE)

        print(
            f"Planilha atualizada: linha {numero_linha}, "
            f"CPF {cpf_registro}, PRONTO = {status}"
        )

    finally:
        workbook.close()

def atualizar_status_planilha(registro, pronto, status_anexo):
    """
    Atualiza PRONTO e STATUS ANEXO do candidato correto.
    PRONTO: OK ou NÃO
    STATUS ANEXO: SIM ou NÃO
    """
    if pronto not in {"OK", "NÃO"}:
        raise ValueError(f"Status PRONTO inválido: {pronto!r}")

    if status_anexo not in {"SIM", "NÃO"}:
        raise ValueError(f"Status ANEXO inválido: {status_anexo!r}")

    numero_linha = registro.get("LINHA_PLANILHA")

    if not isinstance(numero_linha, int) or numero_linha < 2:
        raise ValueError(
            f"Linha da planilha inválida: {numero_linha!r}"
        )

    if not ARQUIVO_CONTROLE.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {ARQUIVO_CONTROLE}"
        )

    workbook = openpyxl.load_workbook(ARQUIVO_CONTROLE)

    try:
        if NOME_ABA not in workbook.sheetnames:
            raise ValueError(f"Aba '{NOME_ABA}' não encontrada.")

        planilha = workbook[NOME_ABA]
        colunas = {}

        for coluna in range(1, planilha.max_column + 1):
            cabecalho = normalizar_texto(
                planilha.cell(row=1, column=coluna).value
            )
            if cabecalho:
                colunas[cabecalho] = coluna

        coluna_pronto = colunas.get("PRONTO")
        coluna_anexo = colunas.get("STATUSANEXO")
        coluna_cpf = colunas.get("CPF")

        if not coluna_pronto:
            raise ValueError("Coluna PRONTO não encontrada.")

        if not coluna_anexo:
            raise ValueError("Coluna STATUS ANEXO não encontrada.")

        if not coluna_cpf:
            raise ValueError("Coluna CPF não encontrada.")

        cpf_planilha = re.sub(
            r"\D",
            "",
            str(
                planilha.cell(
                    row=numero_linha,
                    column=coluna_cpf
                ).value or ""
            )
        )

        cpf_registro = re.sub(
            r"\D",
            "",
            str(registro.get("CPF") or "")
        )

        if not cpf_registro or cpf_planilha != cpf_registro:
            raise RuntimeError(
                f"CPF não confere na linha {numero_linha}. "
                f"Esperado: {cpf_registro}; "
                f"encontrado: {cpf_planilha}. "
                "A planilha não foi alterada."
            )

        planilha.cell(
            row=numero_linha,
            column=coluna_pronto
        ).value = pronto

        planilha.cell(
            row=numero_linha,
            column=coluna_anexo
        ).value = status_anexo

        workbook.save(ARQUIVO_CONTROLE)

        print(
            f"Planilha atualizada — linha {numero_linha} | "
            f"CPF {cpf_registro} | "
            f"PRONTO = {pronto} | "
            f"STATUS ANEXO = {status_anexo}"
        )

    finally:
        workbook.close()



def consultar_status_planilha(registro):
    """
    Consulta PRONTO e STATUS ANEXO do candidato.
    Não altera a planilha.

    Retorna:
        ("OK", "SIM")  -> cadastro e anexo concluídos
        ("OK", "NAO")  -> cadastro concluído, anexo pendente
        ("NAO", "NAO") -> precisa processar o candidato
    """

    numero_linha = registro.get("LINHA_PLANILHA")

    if not isinstance(numero_linha, int) or numero_linha < 2:
        raise ValueError(
            f"Linha da planilha inválida: {numero_linha!r}"
        )

    if not ARQUIVO_CONTROLE.exists():
        raise FileNotFoundError(
            f"Planilha não encontrada: {ARQUIVO_CONTROLE}"
        )

    workbook = openpyxl.load_workbook(
        ARQUIVO_CONTROLE,
        read_only=True,
        data_only=True
    )

    try:
        if NOME_ABA not in workbook.sheetnames:
            raise ValueError(
                f"Aba '{NOME_ABA}' não encontrada."
            )

        planilha = workbook[NOME_ABA]

        # Normaliza os cabeçalhos: STATUS ANEXO vira STATUSANEXO.
        colunas = {}

        for coluna in range(1, planilha.max_column + 1):
            cabecalho = normalizar_texto(
                planilha.cell(row=1, column=coluna).value
            )

            if cabecalho:
                colunas[cabecalho] = coluna

        coluna_pronto = colunas.get("PRONTO")
        coluna_anexo = colunas.get("STATUSANEXO")
        coluna_cpf = colunas.get("CPF")

        faltantes = []

        if not coluna_pronto:
            faltantes.append("PRONTO")

        if not coluna_anexo:
            faltantes.append("STATUS ANEXO")

        if not coluna_cpf:
            faltantes.append("CPF")

        if faltantes:
            raise ValueError(
                f"Colunas não encontradas: {faltantes}. "
                f"Cabeçalhos normalizados: {list(colunas.keys())}"
            )

        # Confere o CPF para não consultar a linha errada.
        cpf_planilha = "".join(
            c for c in str(
                planilha.cell(
                    row=numero_linha,
                    column=coluna_cpf
                ).value or ""
            )
            if c.isdigit()
        )

        cpf_registro = "".join(
            c for c in str(registro.get("CPF") or "")
            if c.isdigit()
        )

        if not cpf_registro or cpf_planilha != cpf_registro:
            raise RuntimeError(
                f"CPF não confere na linha {numero_linha}. "
                f"Esperado: {cpf_registro}; "
                f"encontrado: {cpf_planilha}. "
                "A consulta foi interrompida por segurança."
            )

        pronto = normalizar_texto(
            planilha.cell(
                row=numero_linha,
                column=coluna_pronto
            ).value
        )

        anexo = normalizar_texto(
            planilha.cell(
                row=numero_linha,
                column=coluna_anexo
            ).value
        )

        return pronto, anexo

    finally:
        workbook.close()
